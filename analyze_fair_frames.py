"""
分析 Fair frames presentation.pptx 的布局
"""
from pptx import Presentation
import os

def analyze_template(template_path: str):
    """分析模板文件的所有布局"""
    if not os.path.exists(template_path):
        print(f"文件不存在: {template_path}")
        return
    
    prs = Presentation(template_path)
    
    print(f"\n模板文件: {template_path}")
    print(f"总共有 {len(prs.slide_layouts)} 个布局:\n")
    
    for idx, layout in enumerate(prs.slide_layouts):
        print(f"布局 {idx}: {layout.name}")
        
        # 分析占位符
        placeholders = []
        for shape in layout.placeholders:
            placeholder_type = shape.placeholder_format.type
            placeholder_name = shape.name if hasattr(shape, 'name') else f"Placeholder {placeholder_type}"
            placeholders.append(f"{placeholder_name} (Type: {placeholder_type})")
        
        if placeholders:
            print(f"  占位符: {', '.join(placeholders)}")
        print()

if __name__ == "__main__":
    # 分析 Fair frames presentation.pptx
    fair_frames_path = "Fair frames presentation.pptx"
    analyze_template(fair_frames_path)
    
    # 分析原始模板（用于对比）
    original_template = "templates/MasterTemplate.pptx"
    if os.path.exists(original_template):
        print("\n" + "="*60)
        print("原始模板对比:")
        analyze_template(original_template)

